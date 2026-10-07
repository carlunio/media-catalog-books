from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from ...catalog_form_fields import (
    FORM_FIELD_LABELS,
    FORM_FIELD_ORDER,
    IMPORTANT_FIELDS,
    LABEL_COLORS,
    PERSON_NAME_FIELDS,
    READ_ONLY_FIELDS,
    SELECTABLE_FIELDS,
    SELECTABLE_WITH_CUSTOM_VALUE_FIELDS,
    allowed_values_key,
    label_color,
)
from ..database import get_connection
from . import books

WORKBOOK_MAGIC = "media-catalog-books:external-review"
WORKBOOK_FORMAT_VERSION = "1"
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024

DATA_SHEET = "Fichas"
INSTRUCTIONS_SHEET = "Instrucciones"
ORIGINAL_SHEET = "_Original"
LISTS_SHEET = "_Listas"
META_SHEET = "_Metadatos"

TECHNICAL_FIELDS: tuple[str, ...] = ("id", "form_status")
WORKBOOK_FIELDS: tuple[str, ...] = (*TECHNICAL_FIELDS, *FORM_FIELD_ORDER)
WORKBOOK_EDITABLE_FIELDS: tuple[str, ...] = tuple(
    field
    for field in FORM_FIELD_ORDER
    if field not in READ_ONLY_FIELDS and field != "descripcion"
)
WORKBOOK_READ_ONLY_FIELDS: frozenset[str] = frozenset(
    {*READ_ONLY_FIELDS, "descripcion"}
)

FIELD_LABELS: dict[str, str] = {
    "id": "Ref. del artículo",
    "form_status": "Estado de ficha",
    **FORM_FIELD_LABELS,
}

STATUS_LABELS = {
    "draft": "Borrador",
    "consolidated": "Consolidada",
}
STATUS_FILTERS = {"all", *STATUS_LABELS}

INTEGER_FIELDS = frozenset(books.CORE_BOOKS_INT_FIELDS)
DECIMAL_FIELDS = frozenset(books.CORE_BOOKS_DECIMAL_FIELDS)

THIN_GREY = Side(style="thin", color="A6A6A6")
HEADER_FONT = Font(bold=True, color="1F2937")
LOCKED_FILL = PatternFill("solid", fgColor="E7E6E6")
IMPORTANT_FILL = PatternFill("solid", fgColor="F4CCCC")
EDITABLE_FILL = PatternFill("solid", fgColor="FFFFFF")


def _normalize_status_filter(value: str | None) -> str:
    status = str(value or "all").strip().lower()
    if status not in STATUS_FILTERS:
        raise ValueError("Estado de ficha no válido. Usa all, draft o consolidated.")
    return status


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def _same_value(left: Any, right: Any) -> bool:
    return books._core_values_equal(left, right)


def _safe_text_cell(cell: Any, value: Any) -> None:
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _visible_value(field: str, value: Any) -> Any:
    if value is None:
        return None
    if field == "form_status":
        status = str(value).strip().lower()
        return STATUS_LABELS.get(status, status)
    if field in INTEGER_FIELDS:
        return int(value)
    if field in DECIMAL_FIELDS:
        return float(value)
    return str(value)


def _original_value(field: str, value: Any) -> str:
    if value is None:
        return ""
    if field in INTEGER_FIELDS:
        return str(int(value))
    if field in DECIMAL_FIELDS:
        return f"{float(value):.2f}"
    return str(value)


def _original_hash(rows: list[dict[str, str]]) -> str:
    payload = json.dumps(
        rows,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _query_export_records(
    *,
    block: str,
    module: str,
    form_status: str,
) -> list[dict[str, Any]]:
    selected_columns = ", ".join(f"b.{field}" for field in FORM_FIELD_ORDER)
    sql = f"""
        SELECT bi.id,
               COALESCE(bi.form_status, 'draft') AS form_status,
               {selected_columns}
        FROM book_items AS bi
        JOIN books AS b ON b.id = bi.id
        WHERE bi.block = ? AND bi.module = ?
    """
    params: list[Any] = [block, module]
    if form_status != "all":
        sql += " AND COALESCE(bi.form_status, 'draft') = ?"
        params.append(form_status)
    sql += " ORDER BY bi.id"

    with get_connection(read_only=True) as con:
        rows = con.execute(sql, params).fetchall()

    records: list[dict[str, Any]] = []
    for row in rows:
        record = {
            "id": str(row[0] or "").strip(),
            "form_status": str(row[1] or "draft").strip().lower(),
        }
        for index, field in enumerate(FORM_FIELD_ORDER, start=2):
            record[field] = _json_value(row[index])
        records.append(record)
    return records


def _allowed_values_with_current(
    records: list[dict[str, Any]],
) -> dict[str, list[str]]:
    configured = books.get_books_allowed_values()
    output: dict[str, list[str]] = {}

    for field in SELECTABLE_FIELDS:
        key = allowed_values_key(field)
        values = [
            str(item).strip() for item in configured.get(key, []) if str(item).strip()
        ]
        for record in records:
            current = str(record.get(field) or "").strip()
            if current and current not in values:
                values.append(current)
        output[field] = values
    return output


def _add_instructions_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(INSTRUCTIONS_SHEET)
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 4
    sheet.column_dimensions["B"].width = 105

    sheet["B2"] = "Revisión externa de fichas"
    sheet["B2"].font = Font(size=16, bold=True, color="1F4E78")

    instructions = [
        "Edita únicamente las celdas blancas o rojas de la hoja Fichas.",
        "No cambies las referencias, el estado ni las columnas grises.",
        (
            "Los desplegables son obligatorios salvo Categoría y Género, "
            "donde puedes escribir un valor nuevo."
        ),
        (
            "Introduce cantidades y medidas como números enteros. El precio "
            "admite decimales y se conserva como número, con independencia "
            "de que Excel muestre coma o punto."
        ),
        (
            "Escribe los nombres como «Apellidos, Nombre». Separa varias "
            "personas con punto y coma."
        ),
        (
            "La descripción está bloqueada: se volverá a generar "
            "automáticamente al importar cambios."
        ),
        (
            "Devuelve este mismo archivo .xlsx. La aplicación comprobará "
            "primero todos los cambios y pedirá confirmación antes de guardarlos."
        ),
    ]
    for row_number, instruction in enumerate(instructions, start=4):
        sheet.cell(row=row_number, column=1, value=row_number - 3)
        sheet.cell(row=row_number, column=2, value=instruction)
        sheet.cell(row=row_number, column=2).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.row_dimensions[row_number].height = 34


def _add_meta_sheet(
    workbook: Workbook,
    *,
    block: str,
    module: str,
    form_status: str,
    exported_at: str,
    original_sha256: str,
) -> None:
    sheet = workbook.create_sheet(META_SHEET)
    rows = [
        ("magic", WORKBOOK_MAGIC),
        ("format_version", WORKBOOK_FORMAT_VERSION),
        ("block", block),
        ("module", module),
        ("form_status", form_status),
        ("exported_at_utc", exported_at),
        ("original_sha256", original_sha256),
        ("text_format", "Unicode OOXML (.xlsx)"),
        ("description_policy", "regenerate_on_import"),
    ]
    for row in rows:
        sheet.append(row)
    sheet.sheet_state = "veryHidden"


def _add_original_sheet(
    workbook: Workbook,
    original_rows: list[dict[str, str]],
) -> None:
    sheet = workbook.create_sheet(ORIGINAL_SHEET)
    for column, field in enumerate(WORKBOOK_FIELDS, start=1):
        _safe_text_cell(sheet.cell(row=1, column=column), FIELD_LABELS[field])

    for row_number, record in enumerate(original_rows, start=2):
        for column, field in enumerate(WORKBOOK_FIELDS, start=1):
            _safe_text_cell(
                sheet.cell(row=row_number, column=column),
                record.get(field, ""),
            )
    sheet.sheet_state = "veryHidden"


def _add_allowed_values(
    workbook: Workbook,
    *,
    sheet: Any,
    records: list[dict[str, Any]],
) -> None:
    lists_sheet = workbook.create_sheet(LISTS_SHEET)
    allowed_values = _allowed_values_with_current(records)
    last_data_row = max(2, len(records) + 1)

    for list_column, field in enumerate(
        sorted(SELECTABLE_FIELDS, key=FORM_FIELD_ORDER.index),
        start=1,
    ):
        values = allowed_values.get(field, [])
        if not values:
            continue

        _safe_text_cell(
            lists_sheet.cell(row=1, column=list_column),
            FORM_FIELD_LABELS[field],
        )
        for row_number, value in enumerate(values, start=2):
            _safe_text_cell(
                lists_sheet.cell(row=row_number, column=list_column),
                value,
            )

        column_letter = get_column_letter(list_column)
        range_name = f"opciones_{field}"
        attr_text = (
            "'"
            + LISTS_SHEET
            + "'!"
            + "$"
            + column_letter
            + "$2:"
            + "$"
            + column_letter
            + "$"
            + str(len(values) + 1)
        )
        workbook.defined_names.add(DefinedName(range_name, attr_text=attr_text))

        data_column = WORKBOOK_FIELDS.index(field) + 1
        data_column_letter = get_column_letter(data_column)
        accepts_custom = field in SELECTABLE_WITH_CUSTOM_VALUE_FIELDS
        validation = DataValidation(
            type="list",
            formula1=f"={range_name}",
            allow_blank=True,
        )
        validation.showInputMessage = True
        validation.promptTitle = FORM_FIELD_LABELS[field]
        if accepts_custom:
            validation.prompt = "Elige un valor de la lista o escribe uno nuevo."
            validation.showErrorMessage = False
        else:
            validation.prompt = "Elige un valor de la lista."
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            validation.errorTitle = "Valor no permitido"
            validation.error = "Selecciona uno de los valores del desplegable."
        sheet.add_data_validation(validation)
        validation.add(f"{data_column_letter}2:{data_column_letter}{last_data_row}")

    lists_sheet.sheet_state = "veryHidden"


def create_review_workbook(
    *,
    block: str | None,
    module: str | None,
    form_status: str = "draft",
) -> dict[str, Any]:
    scope_block, scope_module = books.resolve_scope(block, module, require=True)
    assert scope_block is not None
    assert scope_module is not None
    normalized_status = _normalize_status_filter(form_status)
    records = _query_export_records(
        block=scope_block,
        module=scope_module,
        form_status=normalized_status,
    )

    original_rows = [
        {field: _original_value(field, record.get(field)) for field in WORKBOOK_FIELDS}
        for record in records
    ]
    original_sha256 = _original_hash(original_rows)
    exported_at = datetime.now(timezone.utc).isoformat(timespec="seconds")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(WORKBOOK_FIELDS))}{max(1, len(records) + 1)}"
    )

    for column, field in enumerate(WORKBOOK_FIELDS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.value = FIELD_LABELS[field]
        if field == "id":
            color = LABEL_COLORS["lbl-blue"]
        elif field == "form_status":
            color = LABEL_COLORS["lbl-purple"]
        else:
            color = label_color(field)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=THIN_GREY, right=THIN_GREY)

        if field in WORKBOOK_READ_ONLY_FIELDS or field in TECHNICAL_FIELDS:
            message = "Columna informativa; no se importa."
        elif field in SELECTABLE_WITH_CUSTOM_VALUE_FIELDS:
            message = "Desplegable ampliable: también admite valores nuevos."
        elif field in SELECTABLE_FIELDS:
            message = "Campo cerrado: usa un valor del desplegable."
        elif field in PERSON_NAME_FIELDS:
            message = "Formato recomendado: Apellidos, Nombre."
        elif field in INTEGER_FIELDS:
            message = "Número entero."
        elif field in DECIMAL_FIELDS:
            message = "Número decimal; Excel puede mostrar coma o punto."
        else:
            message = "Campo editable."
        cell.comment = Comment(message, "Media Catalog Books")

    sheet.row_dimensions[1].height = 42

    for row_number, record in enumerate(records, start=2):
        for column, field in enumerate(WORKBOOK_FIELDS, start=1):
            cell = sheet.cell(row=row_number, column=column)
            _safe_text_cell(cell, _visible_value(field, record.get(field)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GREY, right=THIN_GREY)
            editable = field in WORKBOOK_EDITABLE_FIELDS
            cell.protection = Protection(locked=not editable)
            if not editable:
                cell.fill = LOCKED_FILL
            elif field in IMPORTANT_FIELDS:
                cell.fill = IMPORTANT_FILL
            else:
                cell.fill = EDITABLE_FILL

            if field in DECIMAL_FIELDS:
                cell.number_format = "0.00"
            elif field in INTEGER_FIELDS:
                cell.number_format = "0"
            elif field in {"id", "isbn", "anio"}:
                cell.number_format = "@"

    for column, field in enumerate(WORKBOOK_FIELDS, start=1):
        label_width = len(FIELD_LABELS[field]) + 3
        if field in {"descripcion", "desperfectos", "detalle_encuadernacion"}:
            width = 40
        elif field in {"titulo", "titulo_completo", "obra_completa"}:
            width = 32
        elif field in PERSON_NAME_FIELDS:
            width = 28
        else:
            width = min(max(label_width, 14), 24)
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.protection.sheet = True
    sheet.protection.set_password("revision")
    sheet.protection.autoFilter = False
    sheet.protection.sort = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False

    _add_instructions_sheet(workbook)
    _add_allowed_values(workbook, sheet=sheet, records=records)
    _add_original_sheet(workbook, original_rows)
    _add_meta_sheet(
        workbook,
        block=scope_block,
        module=scope_module,
        form_status=normalized_status,
        exported_at=exported_at,
        original_sha256=original_sha256,
    )

    workbook.active = workbook.sheetnames.index(DATA_SHEET)
    workbook.properties.creator = "Media Catalog Books"
    workbook.properties.title = f"Revisión de fichas {scope_block}/{scope_module}"
    workbook.properties.subject = "Revisión externa de fichas bibliográficas"
    workbook.properties.description = (
        "Libro de trabajo Unicode OOXML para exportar, revisar e importar fichas."
    )
    workbook.properties.language = "es-ES"

    buffer = BytesIO()
    workbook.save(buffer)
    filename = (
        f"revision_{scope_module}{scope_block}_{normalized_status}_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    return {
        "content": buffer.getvalue(),
        "filename": filename,
        "rows": len(records),
        "block": scope_block,
        "module": scope_module,
        "form_status": normalized_status,
    }


def decode_workbook_base64(content_base64: str) -> bytes:
    import base64

    text = str(content_base64 or "").strip()
    if not text:
        raise ValueError("El archivo Excel está vacío.")
    try:
        content = base64.b64decode(text, validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError("El contenido del archivo Excel no es válido.") from exc
    if len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError("El archivo Excel supera el límite de 25 MB.")
    if not content.startswith(b"PK"):
        raise ValueError("El archivo no es un .xlsx válido.")
    return content


def _read_meta(workbook: Any) -> dict[str, str]:
    if META_SHEET not in workbook.sheetnames:
        raise ValueError("Falta la hoja interna de metadatos.")
    sheet = workbook[META_SHEET]
    meta: dict[str, str] = {}
    for key, value in sheet.iter_rows(min_col=1, max_col=2, values_only=True):
        name = str(key or "").strip()
        if name:
            meta[name] = str(value or "").strip()

    if meta.get("magic") != WORKBOOK_MAGIC:
        raise ValueError(
            "El Excel no fue generado por la revisión externa de esta aplicación."
        )
    if meta.get("format_version") != WORKBOOK_FORMAT_VERSION:
        raise ValueError(
            "La versión de la plantilla no es compatible; exporta un Excel nuevo."
        )
    return meta


def _read_rows(sheet: Any) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    expected_headers = [FIELD_LABELS[field] for field in WORKBOOK_FIELDS]
    actual_headers = [
        str(sheet.cell(row=1, column=column).value or "").strip()
        for column in range(1, len(WORKBOOK_FIELDS) + 1)
    ]
    if actual_headers != expected_headers:
        raise ValueError(
            "Las cabeceras del Excel han cambiado; vuelve a exportar la plantilla."
        )

    rows: list[dict[str, Any]] = []
    formulas: list[dict[str, str]] = []
    for row_number in range(2, sheet.max_row + 1):
        cells = [
            sheet.cell(row=row_number, column=column)
            for column in range(1, len(WORKBOOK_FIELDS) + 1)
        ]
        if not any(cell.value not in (None, "") for cell in cells):
            continue
        record: dict[str, Any] = {}
        formula_fields: dict[str, str] = {}
        for field, cell in zip(WORKBOOK_FIELDS, cells):
            record[field] = cell.value
            if cell.data_type == "f":
                formula_fields[field] = str(cell.value or "")
        rows.append(record)
        formulas.append(formula_fields)
    return rows, formulas


def _normalize_integer(field: str, value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        raise ValueError("debe ser un número entero")
    text = str(value).strip().replace(",", ".")
    try:
        decimal = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("debe ser un número entero") from exc
    if decimal != decimal.to_integral_value():
        raise ValueError("debe ser un número entero, sin decimales")
    number = int(decimal)
    if number < 0:
        raise ValueError("no puede ser negativo")
    return number


def _normalize_decimal(field: str, value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        raise ValueError("debe ser un número")
    text = str(value).strip().replace("€", "").replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("debe ser un número decimal") from exc
    if number < 0:
        raise ValueError("no puede ser negativo")
    return float(number.quantize(Decimal("0.01")))


def _normalize_import_value(field: str, value: Any) -> Any:
    if field in INTEGER_FIELDS:
        return _normalize_integer(field, value)
    if field in DECIMAL_FIELDS:
        return _normalize_decimal(field, value)
    if isinstance(value, datetime):
        raise ValueError("debe escribirse como texto")
    return books._normalize_core_input_value(field, value)


def _canonical_allowed_value(
    field: str,
    value: Any,
    allowed_values: dict[str, list[str]],
) -> tuple[Any, bool]:
    if value in (None, ""):
        return value, True
    options = allowed_values.get(allowed_values_key(field), [])
    by_casefold = {
        str(option).strip().casefold(): str(option).strip()
        for option in options
        if str(option).strip()
    }
    canonical = by_casefold.get(str(value).strip().casefold())
    if canonical is not None:
        return canonical, True
    return value, field in SELECTABLE_WITH_CUSTOM_VALUE_FIELDS


def _query_current_records(ids: list[str]) -> dict[str, dict[str, Any]]:
    if not ids:
        return {}
    placeholders = ", ".join(["?"] * len(ids))
    selected_columns = ", ".join(f"b.{field}" for field in FORM_FIELD_ORDER)
    with get_connection(read_only=True) as con:
        rows = con.execute(
            f"""
            SELECT bi.id, bi.block, bi.module,
                   COALESCE(bi.form_status, 'draft'),
                   {selected_columns}
            FROM book_items AS bi
            JOIN books AS b ON b.id = bi.id
            WHERE bi.id IN ({placeholders})
            """,
            ids,
        ).fetchall()

    output: dict[str, dict[str, Any]] = {}
    for row in rows:
        record = {
            "id": str(row[0] or "").strip(),
            "block": str(row[1] or "").strip(),
            "module": str(row[2] or "").strip(),
            "form_status": str(row[3] or "draft").strip().lower(),
        }
        for index, field in enumerate(FORM_FIELD_ORDER, start=4):
            record[field] = _json_value(row[index])
        output[record["id"]] = record
    return output


def _name_warnings(field: str, value: Any) -> list[str]:
    if field not in PERSON_NAME_FIELDS:
        return []
    warnings: list[str] = []
    for item in re.split(r"[;\n]+", str(value or "")):
        name = item.strip()
        if not name or "," in name:
            continue
        if len(name.split()) > 1:
            warnings.append(
                f"«{name}» no sigue el formato recomendado «Apellidos, Nombre»."
            )
    return warnings


def _isbn_warning(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    compact = re.sub(r"[^0-9Xx]", "", text)
    if len(compact) not in {10, 13}:
        return "El ISBN no tiene 10 ni 13 caracteres significativos."
    return None


def _public_report(inspection: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in inspection.items() if not key.startswith("_")}


def _inspect_workbook(content: bytes) -> dict[str, Any]:
    if len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError("El archivo Excel supera el límite de 25 MB.")
    try:
        workbook = load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise ValueError("No se pudo abrir el archivo .xlsx.") from exc

    meta = _read_meta(workbook)
    for sheet_name in (DATA_SHEET, ORIGINAL_SHEET):
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Falta la hoja obligatoria «{sheet_name}».")

    visible_rows, formulas = _read_rows(workbook[DATA_SHEET])
    original_rows_raw, original_formulas = _read_rows(workbook[ORIGINAL_SHEET])
    if any(original_formulas):
        raise ValueError("La hoja interna original contiene fórmulas no válidas.")

    original_rows = [
        {field: _original_value(field, row.get(field)) for field in WORKBOOK_FIELDS}
        for row in original_rows_raw
    ]
    if _original_hash(original_rows) != meta.get("original_sha256"):
        raise ValueError(
            "La copia interna original ha cambiado; exporta una plantilla nueva."
        )

    sha256 = hashlib.sha256(content).hexdigest()
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    changes: list[dict[str, Any]] = []

    def add_error(book_id: str, field: str, message: str) -> None:
        errors.append(
            {
                "id": book_id,
                "campo": FIELD_LABELS.get(field, field),
                "mensaje": message,
            }
        )

    def add_warning(book_id: str, field: str, message: str) -> None:
        warnings.append(
            {
                "id": book_id,
                "campo": FIELD_LABELS.get(field, field),
                "mensaje": message,
            }
        )

    original_by_id = {
        str(row.get("id") or "").strip(): row
        for row in original_rows
        if str(row.get("id") or "").strip()
    }
    visible_by_id: dict[str, dict[str, Any]] = {}
    formulas_by_id: dict[str, dict[str, str]] = {}
    for row, row_formulas in zip(visible_rows, formulas):
        book_id = str(row.get("id") or "").strip()
        if not book_id:
            add_error("", "id", "La referencia no puede quedar vacía.")
            continue
        if book_id in visible_by_id:
            add_error(book_id, "id", "La referencia aparece más de una vez.")
            continue
        visible_by_id[book_id] = row
        formulas_by_id[book_id] = row_formulas

    missing_ids = sorted(set(original_by_id) - set(visible_by_id))
    unexpected_ids = sorted(set(visible_by_id) - set(original_by_id))
    for book_id in missing_ids:
        add_error(book_id, "id", "Se ha eliminado la fila original.")
    for book_id in unexpected_ids:
        add_error(book_id, "id", "La referencia se ha añadido o modificado.")

    block, module = books.resolve_scope(
        meta.get("block"),
        meta.get("module"),
        require=True,
    )
    assert block is not None
    assert module is not None
    current_by_id = _query_current_records(sorted(original_by_id))
    allowed_values = {
        key: list(values) for key, values in books.get_books_allowed_values().items()
    }
    for field in SELECTABLE_FIELDS:
        key = allowed_values_key(field)
        allowed_values.setdefault(key, [])
        for original in original_rows:
            value = str(original.get(field) or "").strip()
            if value and value not in allowed_values[key]:
                allowed_values[key].append(value)

    updates_by_id: dict[str, dict[str, Any]] = {}
    expected_current: dict[str, dict[str, Any]] = {}

    for book_id in sorted(set(original_by_id) & set(visible_by_id)):
        original = original_by_id[book_id]
        visible = visible_by_id[book_id]
        row_formulas = formulas_by_id.get(book_id, {})
        current = current_by_id.get(book_id)
        if current is None:
            add_error(book_id, "id", "La ficha ya no existe en la base de datos.")
            continue
        if current.get("block") != block or current.get("module") != module:
            add_error(
                book_id,
                "id",
                "La ficha ya no pertenece al bloque y módulo exportados.",
            )
            continue

        original_status = str(original.get("form_status") or "").strip().lower()
        expected_status_label = STATUS_LABELS.get(
            original_status,
            original_status,
        )
        visible_status = str(visible.get("form_status") or "").strip()
        if visible_status != expected_status_label:
            add_error(
                book_id,
                "form_status",
                "El estado es informativo y no puede modificarse en Excel.",
            )
        if str(current.get("form_status") or "").strip().lower() != original_status:
            add_error(
                book_id,
                "form_status",
                "El estado cambió en la aplicación después de exportar el Excel.",
            )

        for field in WORKBOOK_READ_ONLY_FIELDS:
            if field in row_formulas:
                add_error(book_id, field, "No se admiten fórmulas.")
                continue
            try:
                visible_value = _normalize_import_value(field, visible.get(field))
                original_value = _normalize_import_value(field, original.get(field))
            except ValueError:
                add_error(
                    book_id,
                    field,
                    "La columna es informativa y no puede modificarse.",
                )
                continue
            if not _same_value(visible_value, original_value):
                add_error(
                    book_id,
                    field,
                    "La columna es informativa y no puede modificarse.",
                )

        user_updates: dict[str, Any] = {}
        original_changed_values: dict[str, Any] = {}
        for field in WORKBOOK_EDITABLE_FIELDS:
            if field in row_formulas:
                add_error(book_id, field, "No se admiten fórmulas.")
                continue
            try:
                candidate = _normalize_import_value(field, visible.get(field))
                original_value = _normalize_import_value(field, original.get(field))
            except ValueError as exc:
                add_error(book_id, field, str(exc))
                continue

            if _same_value(candidate, original_value):
                continue

            if field in SELECTABLE_FIELDS:
                candidate, accepted = _canonical_allowed_value(
                    field,
                    candidate,
                    allowed_values,
                )
                if not accepted:
                    add_error(
                        book_id,
                        field,
                        "El valor no pertenece a la lista cerrada.",
                    )
                    continue

            if not _same_value(current.get(field), original_value):
                add_error(
                    book_id,
                    field,
                    "El valor cambió en la aplicación después de exportar el Excel.",
                )
                continue

            user_updates[field] = candidate
            original_changed_values[field] = original_value
            for warning in _name_warnings(field, candidate):
                add_warning(book_id, field, warning)
            if field == "isbn":
                warning = _isbn_warning(candidate)
                if warning:
                    add_warning(book_id, field, warning)

        if not user_updates:
            continue

        merged = {**current, **user_updates}
        if "isbn" in user_updates or "palabras_clave" in user_updates:
            normalized_keywords = books._normalize_keywords_for_isbn(
                merged.get("palabras_clave"),
                isbn=merged.get("isbn"),
            )
            merged["palabras_clave"] = normalized_keywords
            if not _same_value(
                current.get("palabras_clave"),
                normalized_keywords,
            ):
                user_updates["palabras_clave"] = normalized_keywords

        description = books.build_core_description(merged)
        if not _same_value(current.get("descripcion"), description):
            user_updates["descripcion"] = description

        effective_updates = {
            field: value
            for field, value in user_updates.items()
            if not _same_value(current.get(field), value)
        }
        if not effective_updates:
            continue

        updates_by_id[book_id] = effective_updates
        expected_current[book_id] = {
            "form_status": current.get("form_status"),
            **{field: current.get(field) for field in effective_updates},
        }
        for field, value in effective_updates.items():
            changes.append(
                {
                    "id": book_id,
                    "campo": FIELD_LABELS[field],
                    "anterior": _json_value(current.get(field)),
                    "nuevo": _json_value(value),
                    "origen": (
                        "Automático"
                        if field in {"descripcion", "palabras_clave"}
                        and field not in original_changed_values
                        else "Excel"
                    ),
                }
            )

    changed_books = len(updates_by_id)
    report: dict[str, Any] = {
        "ok": not errors,
        "valid": not errors,
        "can_apply": not errors and changed_books > 0,
        "workbook_sha256": sha256,
        "workbook": {
            "block": block,
            "module": module,
            "form_status": meta.get("form_status"),
            "exported_at_utc": meta.get("exported_at_utc"),
            "format": "Unicode OOXML (.xlsx)",
        },
        "rows": len(original_rows),
        "changed_books": changed_books,
        "changed_fields": len(changes),
        "errors_count": len(errors),
        "warnings_count": len(warnings),
        "errors": errors,
        "warnings": warnings,
        "changes": changes,
        "applied": False,
        "_updates": updates_by_id,
        "_expected_current": expected_current,
    }
    return report


def preview_review_workbook(content: bytes) -> dict[str, Any]:
    return _public_report(_inspect_workbook(content))


def apply_review_workbook(
    content: bytes,
    *,
    expected_sha256: str,
    consolidate_after_import: bool = False,
) -> dict[str, Any]:
    inspection = _inspect_workbook(content)
    report = _public_report(inspection)
    if expected_sha256 != report["workbook_sha256"]:
        report["ok"] = False
        report["valid"] = False
        report["can_apply"] = False
        report["errors_count"] = int(report["errors_count"]) + 1
        report["errors"] = [
            *report["errors"],
            {
                "id": "",
                "campo": "Archivo",
                "mensaje": (
                    "El archivo cambió después de la vista previa; "
                    "vuelve a analizarlo."
                ),
            },
        ]
        return report
    if not report["can_apply"]:
        return report

    updates_by_id = inspection["_updates"]
    expected_current = inspection["_expected_current"]

    with get_connection() as con:
        con.execute("BEGIN TRANSACTION")
        try:
            for book_id, updates in updates_by_id.items():
                guard = expected_current[book_id]
                fields = list(updates)
                selected = ", ".join(f"b.{field}" for field in fields)
                row = con.execute(
                    f"""
                    SELECT COALESCE(bi.form_status, 'draft'), {selected}
                    FROM book_items AS bi
                    JOIN books AS b ON b.id = bi.id
                    WHERE bi.id = ?
                    """,
                    [book_id],
                ).fetchone()
                if row is None:
                    raise ValueError(
                        f"La ficha {book_id} ya no existe en la base de datos."
                    )
                if not _same_value(row[0], guard.get("form_status")):
                    raise ValueError(
                        f"El estado de {book_id} cambió después de la vista previa."
                    )
                for index, field in enumerate(fields, start=1):
                    if not _same_value(row[index], guard.get(field)):
                        raise ValueError(
                            f"El campo «{FIELD_LABELS[field]}» de {book_id} "
                            "cambió después de la vista previa."
                        )

                assignments = ", ".join(f"{field} = ?" for field in fields)
                con.execute(
                    f"UPDATE books SET {assignments} WHERE id = ?",
                    [*[updates[field] for field in fields], book_id],
                )

                if consolidate_after_import:
                    con.execute(
                        """
                        UPDATE book_items
                        SET form_status = 'consolidated',
                            form_consolidated_at = CURRENT_TIMESTAMP,
                            workflow_status = 'done',
                            workflow_current_node = 'form_consolidated',
                            workflow_action = NULL,
                            workflow_needs_review = FALSE,
                            workflow_review_reason = NULL,
                            pipeline_stage = 'done',
                            updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?
                        """,
                        [book_id],
                    )
                else:
                    con.execute(
                        """
                        UPDATE book_items
                        SET updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?
                        """,
                        [book_id],
                    )
            con.execute("COMMIT")
        except Exception as exc:
            con.execute("ROLLBACK")
            report["ok"] = False
            report["valid"] = False
            report["can_apply"] = False
            report["errors_count"] = int(report["errors_count"]) + 1
            report["errors"] = [
                *report["errors"],
                {
                    "id": "",
                    "campo": "Base de datos",
                    "mensaje": str(exc),
                },
            ]
            return report

    report["ok"] = True
    report["applied"] = True
    report["can_apply"] = False
    report["applied_books"] = len(updates_by_id)
    report["consolidated_after_import"] = bool(consolidate_after_import)
    return report
